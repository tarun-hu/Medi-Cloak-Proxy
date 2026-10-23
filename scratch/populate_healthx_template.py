import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_document():
    # Load from original backup template to preserve template metadata, headers, footers, styles
    doc = docx.Document(r'assets\HEALTHX_2026_Problem_Statement_Submission_Template_ORIGINAL_BACKUP.docx')

    # Remove all default body paragraphs while preserving sectPr (page setup, headers, footers)
    body = doc._body._element
    for p in list(body):
        if p.tag.endswith('p'):
            body.remove(p)

    # Color Palette for Executive Presentation
    COLOR_PRIMARY = RGBColor(0x0F, 0x29, 0x4A)     # Deep Midnight Navy
    COLOR_HEADING = RGBColor(0x1E, 0x3A, 0x8A)     # Royal Blue
    COLOR_SUBHEADING = RGBColor(0x1F, 0x29, 0x37)  # Dark Charcoal
    COLOR_TEXT = RGBColor(0x27, 0x27, 0x2A)        # Off-black for optimal reading contrast
    COLOR_MUTED = RGBColor(0x52, 0x52, 0x5B)       # Slate Gray

    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(22)
        run.bold = True
        run.font.color.rgb = COLOR_PRIMARY
        return p

    def add_subtitle(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(14)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(14)
        run.bold = True
        run.font.color.rgb = COLOR_MUTED
        return p

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        run.bold = True
        run.font.color.rgb = COLOR_PRIMARY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(11)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11.5)
        run.bold = True
        run.font.color.rgb = COLOR_HEADING
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(10.5)
        run.bold = True
        run.font.color.rgb = COLOR_SUBHEADING
        return p

    def add_body_p(text, space_after=4, line_spacing=1.15):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(10)
        run.font.color.rgb = COLOR_TEXT
        return p

    def add_field_line(label, value, is_bold_value=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        
        r_lbl = p.add_run(label + ' ')
        r_lbl.font.name = 'Times New Roman'
        r_lbl.font.size = Pt(10)
        r_lbl.bold = True
        r_lbl.font.color.rgb = COLOR_SUBHEADING
        
        r_val = p.add_run(value)
        r_val.font.name = 'Times New Roman'
        r_val.font.size = Pt(10)
        r_val.bold = is_bold_value
        r_val.font.color.rgb = COLOR_TEXT
        return p

    def add_bullet_item(bold_prefix, text, space_after=3.5, indent=0.25):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(indent)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        
        r_bullet = p.add_run('• ')
        r_bullet.font.name = 'Times New Roman'
        r_bullet.font.size = Pt(10)
        r_bullet.bold = True
        r_bullet.font.color.rgb = COLOR_HEADING

        if bold_prefix:
            r_pre = p.add_run(bold_prefix + ' ')
            r_pre.font.name = 'Times New Roman'
            r_pre.font.size = Pt(10)
            r_pre.bold = True
            r_pre.font.color.rgb = COLOR_SUBHEADING

        r_txt = p.add_run(text)
        r_txt.font.name = 'Times New Roman'
        r_txt.font.size = Pt(10)
        r_txt.font.color.rgb = COLOR_TEXT
        return p

    def add_numbered_item(num_str, bold_prefix, text, space_after=3.5, indent=0.25):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(indent)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        
        r_num = p.add_run(num_str + ' ')
        r_num.font.name = 'Times New Roman'
        r_num.font.size = Pt(10)
        r_num.bold = True
        r_num.font.color.rgb = COLOR_HEADING

        if bold_prefix:
            r_pre = p.add_run(bold_prefix + ' ')
            r_pre.font.name = 'Times New Roman'
            r_pre.font.size = Pt(10)
            r_pre.bold = True
            r_pre.font.color.rgb = COLOR_SUBHEADING

        r_txt = p.add_run(text)
        r_txt.font.name = 'Times New Roman'
        r_txt.font.size = Pt(10)
        r_txt.font.color.rgb = COLOR_TEXT
        return p

    # --- 1. Title & Header ---
    add_title('HEALTHX 2026')
    add_subtitle('Problem Statement')

    # --- 2. Section 1: Submission details ---
    add_h1('1. Submission details')
    add_field_line('Problem Statement ID:', '[To be allotted by the organising committee]')
    add_field_line('Problem statement title:', 'MediCloak: Real-Time Privacy Shield & Governance Gateway for Healthcare AI', is_bold_value=True)
    add_field_line('Organisation and department:', 'K.R. Mangalam University - School of Engineering and Technology')
    add_field_line('Submitted by:', 'Tarun (Team Lead), Shreya Chauhan, Vardaan Saini, Ananya Atthrva')
    add_field_line('Email and mobile:', 'tarunnokwal1@gmail.com | +91-8307011215')
    add_field_line('Submission date:', '19th October 2026')

    # Themes & Categories with clean checkbox formatting
    p_theme = doc.add_paragraph()
    p_theme.paragraph_format.space_before = Pt(1)
    p_theme.paragraph_format.space_after = Pt(3)
    p_theme.paragraph_format.line_spacing = 1.15
    r = p_theme.add_run('Primary theme: ')
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10)
    
    r = p_theme.add_run('☐ Sustainable Healthcare     ')
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10)
    r = p_theme.add_run('☑ Digital Health     ')
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10)
    r = p_theme.add_run('☑ AI in Healthcare')
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10)

    p_cat = doc.add_paragraph()
    p_cat.paragraph_format.space_before = Pt(1)
    p_cat.paragraph_format.space_after = Pt(8)
    p_cat.paragraph_format.line_spacing = 1.15
    r = p_cat.add_run('Category: ')
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10)
    
    r = p_cat.add_run('☑ Software     ')
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10)
    r = p_cat.add_run('☐ Hardware     ')
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10)
    r = p_cat.add_run('☐ Hybrid     ')
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10)
    r = p_cat.add_run('☐ Other, specify')
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10)

    # --- 3. Section 2: Background and problem description ---
    add_h1('2. Background and problem description')

    # Sub-section: Background and unmet need
    add_h2('Background and unmet need:')
    add_body_p('Doctors and medical staff are increasingly using popular AI tools like ChatGPT, Claude, Google Gemini, and Perplexity to save time. These tools help them quickly summarize long patient records, check drug interactions, and write discharge notes—saving 2 to 3 hours of paperwork every single day.')
    add_body_p('However, there is a huge hidden danger: Patient Privacy Leaks ("Shadow AI"). To get useful answers from AI, doctors often copy and paste real patient clinical notes into these public websites. Without realizing it, they expose sensitive details like patient names, phone numbers, Aadhaar numbers, ABHA IDs, and diagnoses to public commercial AI servers.')
    add_body_p('This accidentally violates strict healthcare laws like HIPAA (US) and India\'s DPDPA (Digital Personal Data Protection Act 2023), which can lead to severe fines, legal liability, and loss of medical licenses. Many hospitals attempt to fix this by outright blocking AI websites on hospital Wi-Fi. But this doesn\'t work—doctors simply bypass the blocks using personal mobile hotspots because AI makes their job so much faster and more efficient.')
    add_body_p('The unmet need: We should not stop doctors from using AI. Instead, we need a smart, invisible privacy shield that automatically removes private patient data before it leaves the hospital, while keeping the AI accurate, fast, and clinically helpful.')

    # Sub-section: Specific problem to be solved
    add_h2('Specific problem to be solved:')
    add_body_p('MediCloak addresses four critical technical and operational bottlenecks in healthcare AI security:')
    add_bullet_item('Stopping Data Leaks in Real Time (In-Flight Interception):', 'Traditional security tools only check logs hours after a file is uploaded (when the data is already leaked). We need a system that catches and cleans prompts in milliseconds before they leave the hospital network.')
    add_bullet_item('The "Smart Medical Understanding" Dilemma (Clinical Idiom Preservation):', 'Basic privacy tools blindly remove words that look like locations or names. For example, if a doctor writes "Patient has acute Madras Eye and uses a Jaipur Foot", basic tools redact them to "[Location] Eye" and "[Location] Foot"! This destroys the medical meaning and causes AI to give wrong medical advice. We need a system smart enough to protect patient identities while keeping legitimate medical conditions intact.')
    add_bullet_item('Dual-Registry Identifier Complexity (Indian & Global Healthcare):', 'Generic security tools only look for US-based formats (like SSNs). In India, healthcare data involves 14-digit ABHA IDs, 12-digit Aadhaar numbers, and local Hospital OPD/MRN numbers. The tool must natively understand both Indian and global healthcare formats.')
    add_bullet_item('Zero Visibility for Hospitals:', 'Hospital IT and governance teams currently have no idea which AI tools staff are using, what kind of data is being shared, or if any major data leak is happening.')

    # Sub-section: Existing approaches and their limitations
    add_h2('Existing approaches and their limitations:')
    add_bullet_item('Blocking AI Websites (Firewalls):', 'Doctors bypass blocks using their phones and mobile hotspots. Banning AI creates an unmonitored "Shadow AI" culture and slows down medical work.')
    add_bullet_item('Standard Enterprise DLP (Data Loss Prevention):', 'Very slow, requires installing heavy software on every computer, and cannot inspect modern live streaming chat windows (WebSockets) used by ChatGPT and Claude.')
    add_bullet_item('Third-Party Cloud Anonymizers (AWS/Google Cloud DLP):', 'Sending raw patient data to an external cloud service to "clean it" is itself a privacy violation, and it adds annoying lag (200-500ms delay) to every prompt.')
    add_bullet_item('Generic Anonymization Tools:', 'They lack medical context. They destroy vital medical terms (like Madras Eye, Delhi Belly, or Bombay Blood Group) by mistaking them for patient locations.')

    # Sub-section: Proposed Solution: MediCloak
    add_h2('Proposed Solution: MediCloak (And How It Is Better Than Existing Approaches):')
    add_body_p('MediCloak is an intelligent, in-flight privacy gateway and governance system built specifically for healthcare environments. It sits transparently between hospital workstations and external AI services.')
    
    add_h3('How MediCloak Works:')
    add_numbered_item('1.', 'Zero-Touch In-Flight Proxy:', 'When a doctor types a prompt into ChatGPT, Claude, Gemini, or 50+ other supported AI tools, MediCloak catches the outgoing prompt inside the local network before it hits the internet.')
    add_numbered_item('2.', 'Context-Aware Medical Redactor:', 'It instantly strips out patient names, phone numbers, Aadhaar numbers, ABHA IDs, and addresses, replacing them with clean placeholders (e.g., [PATIENT_1], [ABHA_ID_1]).')
    add_numbered_item('3.', 'Clinical Idiom Whitelisting Engine:', 'Unlike dumb tools, MediCloak recognizes medical conditions with place names (Madras Eye, Jaipur Foot, Delhi Belly, Bombay Blood Group) and leaves them intact. The AI receives 100% of the medical context needed to give an accurate diagnosis, while 0% of the patient’s identity is exposed.')
    add_numbered_item('4.', 'Live SOC Audit Dashboard:', 'Hospital IT teams get a real-time cybersecurity dashboard showing active AI connections, prevented data leaks, threat heatmaps, and a side-by-side prompt comparison (Original vs. Sanitized).')
    add_numbered_item('5.', 'Instant Threat Notification:', 'High-severity privacy exposures trigger instant alerts on the security console with risk scoring and mitigation workflows.')

    add_h3('Why MediCloak is Better Than Existing Approaches:')
    add_bullet_item('100% Frictionless for Doctors:', 'Zero software, browser extensions, or plugins required on the doctor\'s computer. They use normal web browsers naturally.')
    add_bullet_item('Sub-15ms Latency:', 'Redaction happens locally inside the hospital network at lightning speed, with no noticeable delay for the user.')
    add_bullet_item('Preserves Clinical Accuracy:', 'By preserving medical idioms, it prevents AI medical hallucinations and wrong clinical suggestions.')
    add_bullet_item('Complete On-Premises Privacy:', 'Patient data never touches any third-party cloud for de-identification. Everything stays within the hospital perimeter.')
    add_bullet_item('Proactive Prevention vs. Passive Logging:', 'It doesn\'t just record a breach after it happens—it blocks and sanitizes the breach in-flight before data leaves the hospital.')

    # Sub-section: Target users and deployment setting
    add_h2('Target users and deployment setting:')
    add_h3('Target Users:')
    add_bullet_item('Doctors, Nurses & Medical Interns:', 'They can safely use top-tier AI tools (ChatGPT, Claude, Gemini, Glass Health) in their normal web browser without fear of leaking patient information or breaking the law.')
    add_bullet_item('Hospital Management & Security Teams (CISOs):', 'They get a live cybersecurity dashboard showing real-time AI usage across the hospital, prevented leaks, and compliance audit logs.')

    add_h3('Deployment Setting:')
    add_bullet_item('Zero Setup for Doctors:', 'Doctors do not need to install any new apps or browser extensions.')
    add_bullet_item('Hospital Gateway Proxy:', 'MediCloak runs as a lightweight network gateway on the hospital server. It automatically monitors and sanitizes traffic for 50+ AI websites in under 15 milliseconds.')
    add_bullet_item('Real-Time SOC Console:', 'High-risk leaks immediately flash on the hospital IT monitoring feed with full audit trails.')

    # Sub-section: Future Scope & Next Steps
    add_h2('Future Scope & Next Steps:')
    add_bullet_item('1. Smart Local Re-hydration:', 'In the next version, MediCloak will locally restore the patient\'s real name when the AI answers back on the doctor\'s screen—so the doctor sees the full name, but the AI cloud never saw it.')
    add_bullet_item('2. Autonomous Voice Escalation (VAPI):', 'Adding an automated AI voice agent (via VAPI) that automatically dials the on-call hospital compliance officer if a severe bulk-data leak attempt is detected.')
    add_bullet_item('3. Hospital Software (EMR/EHR) Integration:', 'Direct connection into hospital software systems (Epic, Cerner, and Indian ABDM/HIMS portals) so doctors can safely summarize records with a single click.')
    add_bullet_item('4. Offline On-Premise AI Models:', 'Running lightweight local AI models on hospital servers for clinics with sensitive data or limited internet.')

    # Sub-section: Notes, Declarations, Signatures (Preserving Template Requirements)
    p_note = doc.add_paragraph()
    p_note.paragraph_format.space_before = Pt(12)
    p_note.paragraph_format.space_after = Pt(4)
    r = p_note.add_run('(')
    r.font.name = 'Times New Roman'
    r.font.size = Pt(9.5)
    r.font.color.rgb = COLOR_MUTED
    r = p_note.add_run('Note: ')
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(9.5)
    r.font.color.rgb = COLOR_MUTED
    r = p_note.add_run('Don’t disclose any confidential information that risk IP (Patent, Industrial design, etc.) Protection)')
    r.font.name = 'Times New Roman'
    r.font.size = Pt(9.5)
    r.font.color.rgb = COLOR_MUTED

    p_decl = doc.add_paragraph()
    p_decl.paragraph_format.space_before = Pt(4)
    p_decl.paragraph_format.space_after = Pt(8)
    p_decl.paragraph_format.line_spacing = 1.15
    r = p_decl.add_run('We confirm originality, acknowledge third-party material, disclose significant AI use, and will not present this concept at another hackathon before HEALTHX 2026 concludes. We will follow event safety and ethics rules.')
    r.font.name = 'Times New Roman'
    r.font.size = Pt(9.5)
    r.font.color.rgb = COLOR_MUTED

    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(10)
    p_sig.paragraph_format.space_after = Pt(6)
    r_sig_label = p_sig.add_run('Team leader signature and date: ')
    r_sig_label.bold = True
    r_sig_label.font.name = 'Times New Roman'
    r_sig_label.font.size = Pt(10.5)
    r_sig_label.font.color.rgb = COLOR_PRIMARY

    r_sig_val = p_sig.add_run('Tarun | 19-10-2026')
    r_sig_val.bold = True
    r_sig_val.font.name = 'Times New Roman'
    r_sig_val.font.size = Pt(10.5)
    r_sig_val.font.color.rgb = COLOR_TEXT

    # Save to target document
    target_path = r'assets\HEALTHX_2026_Problem_Statement_Submission_Template (1).docx'
    doc.save(target_path)
    print(f'Successfully created and formatted: {target_path}')

if __name__ == '__main__':
    create_document()
